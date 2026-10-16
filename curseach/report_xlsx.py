"""
Генерация XLSX-отчёта сканера.
openpyxl — опциональная зависимость: если её нет, is_available() вернёт False.
"""

import datetime
import io
from typing import Any, Dict, List, Optional


# Проверяем, установлен ли openpyxl, но сами модули не импортируем на верхнем уровне.
# Импорты — внутри generate_*(), чтобы Pylance не ругался на "possibly unbound".
_AVAILABLE = False
try:
    import openpyxl  # noqa: F401
    _AVAILABLE = True
except ImportError:
    pass


_LEVEL_RU = {
    "CRITICAL": "Критично",
    "HIGH":     "Высоко",
    "MEDIUM":   "Средне",
    "LOW":      "Низко",
}

_LEVEL_ORDER = {"CRITICAL": 0, "HIGH": 1, "MEDIUM": 2, "LOW": 3}

_LEVEL_COLORS = {
    "CRITICAL": ("FBE9E7", "B3261E"),
    "HIGH":     ("F4E6F8", "7A1FA2"),
    "MEDIUM":   ("FDF0DC", "B87300"),
    "LOW":      ("E6F3EA", "1B6E2F"),
}

_MODE_TITLES = {
    "quick":  "Быстрая проверка",
    "full":   "Полная проверка",
    "custom": "Выборочная проверка",
}


def is_available() -> bool:
    """True, если openpyxl установлен и можно генерировать XLSX."""
    return _AVAILABLE


def calculate_score(critical: int, high: int, medium: int, low: int) -> int:
    penalty = critical * 20 + high * 10 + medium * 5 + low * 2
    return max(0, 100 - penalty)


def generate_xlsx(
    target_url: str,
    mode: str,
    scan_dt: datetime.datetime,
    score: int,
    critical: int,
    high: int,
    medium: int,
    low: int,
    findings: List[Dict[str, Any]],
) -> Optional[bytes]:
    """
    Формирует XLSX-отчёт по одному скану.
    Возвращает bytes или None, если openpyxl недоступен.
    """
    if not _AVAILABLE:
        return None

    from openpyxl import Workbook
    from openpyxl.styles import (
        Alignment,
        Border,
        Font,
        PatternFill,
        Side,
    )

    wb = Workbook()
    ws = wb.active
    if ws is None:
        return None

    # ---- Стили ----
    title_font = Font(name="Calibri", bold=True, size=14, color="111111")
    h1_font = Font(name="Calibri", bold=True, size=12, color="111111")
    bold_font = Font(name="Calibri", bold=True)
    header_font = Font(name="Calibri", bold=True, color="FFFFFF")
    header_fill = PatternFill(fill_type="solid", fgColor="1A1A1A")

    center = Alignment(horizontal="center", vertical="center")
    left_wrap = Alignment(horizontal="left", vertical="top", wrap_text=True)
    left_top = Alignment(horizontal="left", vertical="top")

    thin = Side(border_style="thin", color="D0D4D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    scan_str = scan_dt.strftime("%Y-%m-%d %H:%M:%S")

    # ---- Заголовок ----
    ws.merge_cells("A1:F1")
    cell_a1 = ws["A1"]
    cell_a1.value = "Отчёт об аудите безопасности API"
    cell_a1.font = title_font
    cell_a1.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 24

    ws.merge_cells("A2:F2")
    cell_a2 = ws["A2"]
    cell_a2.value = f"Сформирован: {scan_str}   |   Система: DevSecOps Scanner"
    cell_a2.font = Font(name="Calibri", size=10, color="666666")
    cell_a2.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[2].height = 18

    ws.append([])

    # ---- Метаданные ----
    meta_rows = [
        ("Целевой ресурс", target_url),
        ("Режим сканирования", _MODE_TITLES.get(mode, "Полная проверка")),
        ("Дата и время", scan_str),
        ("Оценка безопасности", f"{score} / 100"),
    ]
    for key, value in meta_rows:
        ws.append([key, value])
        row = ws.max_row
        c1 = ws.cell(row=row, column=1)
        c1.font = bold_font
        c1.alignment = left_top
        c2 = ws.cell(row=row, column=2)
        c2.alignment = left_top
        ws.merge_cells(
            start_row=row, start_column=2, end_row=row, end_column=6,
        )

    ws.append([])

    # ---- Сводка ----
    ws.append(["СВОДКА"])
    ws.cell(row=ws.max_row, column=1).font = h1_font

    ws.append(["Уровень", "Количество"])
    sum_header_row = ws.max_row
    for col in (1, 2):
        c = ws.cell(row=sum_header_row, column=col)
        c.font = header_font
        c.fill = header_fill
        c.alignment = center
        c.border = border

    summary_data = [
        ("CRITICAL", critical),
        ("HIGH", high),
        ("MEDIUM", medium),
        ("LOW", low),
    ]
    for level_key, count in summary_data:
        ws.append([_LEVEL_RU[level_key], count])
        row = ws.max_row
        c1 = ws.cell(row=row, column=1)
        c2 = ws.cell(row=row, column=2)
        c1.border = border
        c2.border = border
        c2.alignment = center
        bg, fg = _LEVEL_COLORS[level_key]
        c1.fill = PatternFill(fill_type="solid", fgColor=bg)
        c1.font = Font(name="Calibri", bold=True, color=fg)

    total = critical + high + medium + low
    ws.append(["Всего", total])
    row = ws.max_row
    c1 = ws.cell(row=row, column=1)
    c2 = ws.cell(row=row, column=2)
    c1.font = bold_font
    c2.font = bold_font
    c1.border = border
    c2.border = border
    c2.alignment = center

    ws.append([])

    # ---- Таблица находок ----
    ws.append(["НАЙДЕННЫЕ УЯЗВИМОСТИ"])
    ws.cell(row=ws.max_row, column=1).font = h1_font

    headers = ["№", "Уровень", "Тип", "Эндпоинт", "Описание", "Рекомендация"]
    ws.append(headers)
    tbl_header_row = ws.max_row
    for col_idx in range(1, len(headers) + 1):
        c = ws.cell(row=tbl_header_row, column=col_idx)
        c.font = header_font
        c.fill = header_fill
        c.alignment = center
        c.border = border

    sorted_findings = sorted(
        findings,
        key=lambda f: (
            _LEVEL_ORDER.get(f.get("level", "LOW"), 99),
            f.get("type", ""),
            f.get("endpoint", ""),
        ),
    )

    if not sorted_findings:
        ws.append(["-", "-", "-", "-", "Уязвимостей не обнаружено", "-"])
        for col_idx in range(1, len(headers) + 1):
            ws.cell(row=ws.max_row, column=col_idx).border = border
    else:
        for idx, f in enumerate(sorted_findings, start=1):
            details = f.get("details") or {}
            level = f.get("level", "")
            ws.append([
                idx,
                _LEVEL_RU.get(level, level),
                f.get("type", ""),
                f.get("endpoint", ""),
                f.get("desc", ""),
                details.get("recommendation", ""),
            ])
            row = ws.max_row
            for col_idx in range(1, len(headers) + 1):
                c = ws.cell(row=row, column=col_idx)
                c.border = border
                c.alignment = left_wrap
            c1 = ws.cell(row=row, column=1)
            c2 = ws.cell(row=row, column=2)
            c1.alignment = Alignment(horizontal="left", vertical="center")
            c2.alignment = center
            if level in _LEVEL_COLORS:
                bg, fg = _LEVEL_COLORS[level]
                c2.fill = PatternFill(fill_type="solid", fgColor=bg)
                c2.font = Font(name="Calibri", bold=True, color=fg)

    # ---- Ширина колонок ----
    fixed_widths = {
        "A": 28,
        "B": 22,
        "C": 26,
        "D": 32,
        "E": 60,
        "F": 55,
    }
    for col_letter, width in fixed_widths.items():
        ws.column_dimensions[col_letter].width = width

    # ---- Заморозка шапки + автофильтр ----
    ws.freeze_panes = f"A{tbl_header_row + 1}"
    ws.auto_filter.ref = f"A{tbl_header_row}:F{ws.max_row}"

    # ---- Сохранение ----
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate_history_xlsx(scans: List[Dict[str, Any]]) -> Optional[bytes]:
    """
    Сводный XLSX по всей истории сканирований.
    Колонки: ID / Дата / Цель / Оценка / Критично / Высоко / Средне / Низко / Всего.
    Данные идут как пришли (обычно — по убыванию id, как в UI истории).
    """
    if not _AVAILABLE:
        return None

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    if ws is None:
        return None
    ws.title = "История"

    title_font = Font(name="Calibri", bold=True, size=14, color="111111")
    bold_font = Font(name="Calibri", bold=True)
    header_font = Font(name="Calibri", bold=True, color="FFFFFF")
    header_fill = PatternFill(fill_type="solid", fgColor="1A1A1A")
    center = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left", vertical="center")

    thin = Side(border_style="thin", color="D0D4D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # ---- Заголовок ----
    ws.merge_cells("A1:I1")
    ws["A1"].value = "История сканирований безопасности API"
    ws["A1"].font = title_font
    ws["A1"].alignment = left
    ws.row_dimensions[1].height = 24

    ws.merge_cells("A2:I2")
    ws["A2"].value = (
        f"Всего записей: {len(scans)}   |   "
        f"Система: DevSecOps Scanner   |   "
        f"Сформирован: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    )
    ws["A2"].font = Font(name="Calibri", size=10, color="666666")
    ws["A2"].alignment = left
    ws.row_dimensions[2].height = 18

    ws.append([])

    # ---- Шапка таблицы ----
    headers = [
        "ID", "Дата", "Цель", "Оценка",
        "Критично", "Высоко", "Средне", "Низко", "Всего",
    ]
    ws.append(headers)
    header_row = ws.max_row
    for col_idx in range(1, len(headers) + 1):
        c = ws.cell(row=header_row, column=col_idx)
        c.font = header_font
        c.fill = header_fill
        c.alignment = center
        c.border = border

    # ---- Строки ----
    if not scans:
        ws.append(["-"] * len(headers))
        for col_idx in range(1, len(headers) + 1):
            ws.cell(row=ws.max_row, column=col_idx).border = border
    else:
        for s in scans:
            crit = s.get("critical", 0) or 0
            high = s.get("high", 0) or 0
            med = s.get("medium", 0) or 0
            low = s.get("low", 0) or 0
            total = crit + high + med + low
            score = calculate_score(crit, high, med, low)

            date_str = (s.get("scan_date") or "").replace("T", " ")[:19]
            ws.append([
                s.get("id", ""),
                date_str,
                s.get("target_url", ""),
                score,
                crit, high, med, low, total,
            ])
            row = ws.max_row
            for col_idx in range(1, len(headers) + 1):
                c = ws.cell(row=row, column=col_idx)
                c.border = border
            # Числовые колонки по центру
            for col_idx in (1, 4, 5, 6, 7, 8, 9):
                ws.cell(row=row, column=col_idx).alignment = center
            # Цвет оценки
            score_cell = ws.cell(row=row, column=4)
            if score >= 80:
                score_cell.font = Font(name="Calibri", bold=True, color="1B6E2F")
            elif score >= 50:
                score_cell.font = Font(name="Calibri", bold=True, color="B87300")
            elif score >= 20:
                score_cell.font = Font(name="Calibri", bold=True, color="7A1FA2")
            else:
                score_cell.font = Font(name="Calibri", bold=True, color="B3261E")

    # ---- Ширина колонок ----
    widths = {
        "A": 8,   # ID
        "B": 20,  # Дата
        "C": 42,  # Цель
        "D": 10,  # Оценка
        "E": 12,  # Критично
        "F": 10,  # Высоко
        "G": 10,  # Средне
        "H": 8,   # Низко
        "I": 8,   # Всего
    }
    for col_letter, width in widths.items():
        ws.column_dimensions[col_letter].width = width

    # ---- Автофильтр по шапке ----
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(headers))}{ws.max_row}"
    ws.freeze_panes = f"A{header_row + 1}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()